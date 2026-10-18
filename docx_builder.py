"""تحويل البحث المكتوب إلى ملف Word منسّق (يدعم العربية من اليمين لليسار)."""
import os
import re
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

PPR_AFTER_BIDI = ("w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
                  "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
                  "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
                  "w:pPrChange")
# في الفقرات ثنائية الاتجاه (bidi) يفسّر Word المحاذاة "left" على أنها بداية السطر، أي اليمين في العربية
START = WD_ALIGN_PARAGRAPH.LEFT
RPR_AFTER_RTL = ("w:cs", "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath")

LABELS = {
    "ar": {"by": "إعداد الطالب", "sup": "إشراف", "spec": "التخصص", "toc": "المحتويات",
           "abstract": "الملخص", "intro": "المقدمة", "chapter": "الفصل", "conc": "الخاتمة",
           "refs": "المصادر والمراجع", "year": "العام الدراسي",
           "ord": ["الأول", "الثاني", "الثالث", "الرابع", "الخامس", "السادس"]},
    "en": {"by": "Prepared by", "sup": "Supervised by", "spec": "Specialization", "toc": "Contents",
           "abstract": "Abstract", "intro": "Introduction", "chapter": "Chapter", "conc": "Conclusion",
           "refs": "References", "year": "Academic Year",
           "ord": ["One", "Two", "Three", "Four", "Five", "Six"]},
}


class Builder:
    def __init__(self, lang: str):
        self.rtl = lang == "ar"
        self.L = LABELS["ar" if self.rtl else "en"]
        self.font = "Simplified Arabic" if self.rtl else "Times New Roman"
        self.doc = Document()
        self._setup()

    # ---------- إعداد الصفحة والأنماط ----------
    def _setup(self):
        sec = self.doc.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.top_margin = sec.bottom_margin = Cm(2.5)
        sec.left_margin = sec.right_margin = Cm(2.5)
        if self.rtl:
            bidi = OxmlElement("w:bidi")
            sec._sectPr.append(bidi)

        base = 14 if self.rtl else 12
        for name, size, bold in (("Normal", base, False), ("Title", 26, True),
                                 ("Heading 1", base + 6, True), ("Heading 2", base + 3, True),
                                 ("Heading 3", base + 1, True), ("List Bullet", base, False),
                                 ("List Number", base, False)):
            st = self.doc.styles[name]
            st.font.name = self.font
            st.font.size = Pt(size)
            st.font.bold = bold
            st.font.cs_bold = bold
            st.font.color.rgb = RGBColor(0x1F, 0x3D, 0x36) if name.startswith(("Heading", "Title")) \
                else RGBColor(0, 0, 0)
            rpr = st.element.get_or_add_rPr()
            fonts = rpr.get_or_add_rFonts()
            for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                fonts.set(qn(attr), self.font)
            self._set_cs_size(rpr, size)
            if self.rtl:
                self._add_rtl(rpr)
                self._add_bidi(st.element.get_or_add_pPr())
            pf = st.paragraph_format
            if name == "Normal":
                pf.line_spacing = 1.5
                pf.space_after = Pt(6)
            elif name.startswith("Heading"):
                pf.space_before = Pt(18 if name == "Heading 1" else 12)
                pf.space_after = Pt(8)
                pf.keep_with_next = True

        # ترقيم الصفحات في التذييل
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._add_field(p.add_run(), "PAGE")
        sec.different_first_page_header_footer = True  # بلا رقم على الغلاف

    @staticmethod
    def _add_bidi(ppr):
        if ppr.find(qn("w:bidi")) is None:
            el = OxmlElement("w:bidi")
            ppr.insert_element_before(el, *PPR_AFTER_BIDI)

    @staticmethod
    def _add_rtl(rpr):
        if rpr.find(qn("w:rtl")) is None:
            rpr.insert_element_before(OxmlElement("w:rtl"), *RPR_AFTER_RTL)

    @staticmethod
    def _set_cs_size(rpr, size):
        el = rpr.find(qn("w:szCs"))
        if el is None:
            el = OxmlElement("w:szCs")
            rpr.insert_element_before(el, "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd",
                                      "w:fitText", "w:vertAlign", "w:rtl", *RPR_AFTER_RTL)
        el.set(qn("w:val"), str(size * 2))

    @staticmethod
    def _add_field(run, instr):
        for tag, attr in (("w:fldChar", "begin"), ("w:instrText", None), ("w:fldChar", "end")):
            el = OxmlElement(tag)
            if attr:
                el.set(qn("w:fldCharType"), attr)
            else:
                el.set(qn("xml:space"), "preserve")
                el.text = instr
            run._r.append(el)

    # ---------- إضافة المحتوى ----------
    def _align(self, p, align=None):
        if self.rtl:
            self._add_bidi(p._p.get_or_add_pPr())
        p.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
        return p

    def _runs(self, p, text, bold=False, size=None):
        # يدعم **نص غامق**
        for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
            if not part:
                continue
            run = p.add_run(part)
            if bold or i % 2 == 1:
                run.font.bold = run.font.cs_bold = True
            if size:
                run.font.size = Pt(size)
                self._set_cs_size(run._r.get_or_add_rPr(), size)
            if self.rtl:
                self._add_rtl(run._r.get_or_add_rPr())

    def para(self, text, style=None, align=None, bold=False, size=None):
        p = self.doc.add_paragraph(style=style)
        self._align(p, align)
        self._runs(p, text, bold, size)
        return p

    def heading(self, text, level):
        p = self.doc.add_heading(level=level)
        self._align(p, WD_ALIGN_PARAGRAPH.CENTER if level == 1 else START)
        self._runs(p, text)
        return p

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def markdown(self, md: str):
        for raw in md.splitlines():
            line = raw.strip()
            if not line or line in ("---", "***"):
                continue
            if m := re.match(r"^#{1,6}\s+(.*)", line):
                level = min(max(len(line) - len(line.lstrip("#")), 2), 3)
                self.heading(m.group(1).strip("* "), level)
            elif m := re.match(r"^[-*•]\s+(.*)", line):
                self.para(m.group(1), style="List Bullet")
            elif re.match(r"^\d+[.)]\s+", line):
                # الرقم يبقى نصاً حتى لا يستمر الترقيم بين القوائم المختلفة
                self.para(line, align=START)
            else:
                self.para(line)

    # ---------- بناء البحث ----------
    def build(self, r, research: dict, path: str):
        L, C, ord_ = self.L, WD_ALIGN_PARAGRAPH.CENTER, self.L["ord"]

        # الغلاف
        for line in (r["university"], r["college"], r["department"]):
            self.para(line, align=C, bold=True, size=16)
        for _ in range(5):
            self.doc.add_paragraph()
        self.para(research["title"], style="Title", align=C)
        for _ in range(4):
            self.doc.add_paragraph()
        self.para(f"{L['by']}: {r['student_name']}", align=C, bold=True, size=15)
        self.para(f"{L['spec']}: {r['specialization']}", align=C, size=14)
        self.para(f"{L['sup']}: {r['supervisor']}", align=C, bold=True, size=15)
        for _ in range(3):
            self.doc.add_paragraph()
        y = date.today().year
        self.para(f"{L['year']} {y - 1}–{y}" if date.today().month < 9 else f"{L['year']} {y}–{y + 1}",
                  align=C, size=13)
        self.page_break()

        # المحتويات
        def chapter_label(i):
            n = ord_[i] if i < len(ord_) else str(i + 1)
            return f"{L['chapter']} {n}: "

        self.heading(L["toc"], 1)
        toc = [L["abstract"], L["intro"]]
        toc += [chapter_label(i) + ch["title"] for i, ch in enumerate(research["chapters"])]
        toc += [L["conc"], L["refs"]]
        for n, item in enumerate(toc, 1):
            self.para(f"{n}. {item}", align=START)
        self.page_break()

        self.heading(L["abstract"], 1)
        self.markdown(research["abstract"])
        self.page_break()

        self.heading(L["intro"], 1)
        self.markdown(research["introduction"])

        for i, ch in enumerate(research["chapters"]):
            self.page_break()
            self.heading(chapter_label(i) + ch["title"], 1)
            self.markdown(ch["body"])

        self.page_break()
        self.heading(L["conc"], 1)
        self.markdown(research["conclusion"])

        self.page_break()
        self.heading(L["refs"], 1)
        for line in research["references"].splitlines():
            line = re.sub(r"^([-*•]|\d+[.)])\s+", "", line.strip())
            if line:
                p = self.para(line, align=START)
                p.paragraph_format.first_line_indent = Cm(-1)
                p.paragraph_format.left_indent = Cm(1)  # تعليق بارز؛ في bidi تُطبَّق على بداية السطر

        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        self.doc.save(path)
        return path


def build_docx(r, research: dict, path: str) -> str:
    return Builder(r["language"]).build(r, research, path)
